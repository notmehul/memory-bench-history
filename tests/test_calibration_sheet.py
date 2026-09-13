"""G4 rater sheet: blinding, the round trip, and the checks that refuse a
sheet rather than guess at it."""

import json
import sys
from pathlib import Path

import pytest

pytest.importorskip("openpyxl")
from openpyxl import load_workbook  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import calibration_sheet as cs  # noqa: E402

PACKET = [
    {"id": "pa0000000001", "criterion": "The note names a responsible person.",
     "output": "Owner: Alice.\n\nDue before the next release."},
    {"id": "pb0000000002", "criterion": "The note does not present the 48h SLA as current.",
     "output": "We used to run a 48h SLA; that was replaced last quarter."},
    {"id": "pc0000000003", "criterion": "The checklist sets a 200-line cap on PRs.",
     "output": "x" * 9000},          # forces the row-height cap
]


@pytest.fixture
def packet_dir(tmp_path: Path) -> Path:
    d = tmp_path / "calibration"
    d.mkdir()
    (d / "rater-packet.json").write_text(json.dumps(PACKET))
    # the answer key lives beside the packet and must never reach the sheet
    (d / "packet-key.json").write_text(json.dumps(
        {"pa0000000001": {"org": "org-00001", "run_id": "P-0007-01:ceiling",
                          "assertion_id": "asrt-1"}}))
    return d


@pytest.fixture
def sheet(tmp_path: Path, packet_dir: Path) -> Path:
    out = tmp_path / "packet.xlsx"
    assert cs.main(["export", str(packet_dir), str(out)]) == 0
    return out


def _fill(path: Path, answers: list, out: Path) -> Path:
    wb = load_workbook(path)
    ws = wb["Rate"]
    for offset, value in enumerate(answers):
        ws.cell(2 + offset, 2, value)
    wb.save(out)
    return out


def test_export_shape(sheet: Path):
    wb = load_workbook(sheet)
    assert wb.sheetnames == ["How to rate", "Rate"]
    ws = wb["Rate"]
    assert ws.max_row == len(PACKET) + 1
    assert ws.freeze_panes == "C2"
    assert [c.value for c in ws[1][:5]] == list(cs.HEADERS)
    assert ws.cell(2, 3).value == PACKET[0]["criterion"]
    assert ws.cell(2, 4).value == PACKET[0]["output"]
    assert ws.cell(2, 5).value == PACKET[0]["id"]


def test_export_has_a_true_false_dropdown(sheet: Path):
    ws = load_workbook(sheet)["Rate"]
    assert any("TRUE,FALSE" in (dv.formula1 or "") for dv in ws.data_validations.dataValidation)


def test_row_height_is_capped(sheet: Path):
    ws = load_workbook(sheet)["Rate"]
    heights = [ws.row_dimensions[i].height for i in range(2, len(PACKET) + 2)]
    assert all(h is not None and h <= cs.MAX_ROW_POINTS for h in heights)
    assert heights[2] == cs.MAX_ROW_POINTS          # the 9k-char output


def test_sheet_stays_blinded(sheet: Path):
    """A rater must not be able to see the run id, condition, or side."""
    blob = json.dumps([[c.value for c in row] for row in load_workbook(sheet)["Rate"].rows])
    for leak in ("P-0007-01", "ceiling", "asrt-1", "org-00001", "twin", "floor"):
        assert leak not in blob, f"{leak!r} leaked into the rater sheet"


def test_round_trip(tmp_path: Path, packet_dir: Path, sheet: Path):
    filled = _fill(sheet, ["TRUE", "FALSE", "TRUE"], tmp_path / "filled.xlsx")
    out = tmp_path / "ratings.json"
    assert cs.main(["import", str(filled), "--packet-dir", str(packet_dir),
                    "--out", str(out)]) == 0
    assert json.loads(out.read_text()) == {
        "pa0000000001": True, "pb0000000002": False, "pc0000000003": True}


def test_round_trip_accepts_common_spellings(tmp_path: Path, packet_dir: Path, sheet: Path):
    filled = _fill(sheet, ["yes", " no ", 1], tmp_path / "filled.xlsx")
    out = tmp_path / "ratings.json"
    assert cs.main(["import", str(filled), "--packet-dir", str(packet_dir),
                    "--out", str(out)]) == 0
    assert list(json.loads(out.read_text()).values()) == [True, False, True]


def test_import_refuses_a_blank_answer(tmp_path: Path, packet_dir: Path, sheet: Path):
    filled = _fill(sheet, ["TRUE", None, "TRUE"], tmp_path / "filled.xlsx")
    with pytest.raises(SystemExit, match="row 3: no answer"):
        cs.main(["import", str(filled), "--packet-dir", str(packet_dir),
                 "--out", str(tmp_path / "r.json")])


def test_import_refuses_a_junk_answer(tmp_path: Path, packet_dir: Path, sheet: Path):
    filled = _fill(sheet, ["TRUE", "maybe", "TRUE"], tmp_path / "filled.xlsx")
    with pytest.raises(SystemExit, match="not TRUE or FALSE"):
        cs.main(["import", str(filled), "--packet-dir", str(packet_dir),
                 "--out", str(tmp_path / "r.json")])


def test_import_refuses_a_mangled_sheet(tmp_path: Path, packet_dir: Path, sheet: Path):
    """Deleting or reordering rows breaks the mapping; fail loudly, never guess."""
    wb = load_workbook(sheet)
    ws = wb["Rate"]
    for offset, value in enumerate(["TRUE", "TRUE", "TRUE"]):
        ws.cell(2 + offset, 2, value)
    ws.delete_rows(3)
    mangled = tmp_path / "mangled.xlsx"
    wb.save(mangled)
    with pytest.raises(SystemExit, match="does not match the packet"):
        cs.main(["import", str(mangled), "--packet-dir", str(packet_dir),
                 "--out", str(tmp_path / "r.json")])
